"""Offline document/provider contracts: no calls to OpenAI, Band, or E2B."""
import asyncio
import io
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from PIL import Image
from sqlmodel import select

from app import documents
from app.openai_provider import LegacyCheckpoint, configuration, migrate_profiles, pin_task
from app.persistence.models import (AgentProfile, AgentTask, Budget, ChannelMember, DirectMessage, Document,
    DocumentCollection, DocumentUsage, Message, OpenAISettings, TaskModel, db)
from app.runtime.tools.documents import document_tools
from app.tasks.usage import over_budget
from helpers import publish, sdk_mock, upload


def test_staged_upload_publish_is_atomic_and_files_only_messages_work(space):
    doc = upload(space)
    url = f"/api/channels/{space.channel.id}/documents"
    assert space.api.get(url, headers=space.auth).json() == []
    message = publish(space, doc)
    assert message["content"] == "" and message["attachments"][0]["status"] == "queued"
    assert len(space.api.get(url, headers=space.auth).json()) == 1
    response = space.api.post(f"/api/channels/{space.channel.id}/messages", headers=space.auth,
        json={"content": "must roll back", "attachment_ids": [doc["id"], "missing"]})
    assert response.status_code == 404
    with db() as session:
        assert not session.exec(select(Message).where(Message.channel_id == space.channel.id,
                                                     Message.content == "must roll back")).first()
    assert space.api.post(f"/api/channels/{space.channel.id}/messages", headers=space.auth,
                          json={"content": "   "}).status_code == 422


def test_document_permissions_include_dm_ownership_and_staging(space):
    doc = upload(space)
    path = f"/api/documents/{doc['id']}"
    assert space.api.get(path + "/download").status_code == 401
    assert space.api.get(path, headers=space.other_auth).status_code == 403
    with db() as session:
        session.add(ChannelMember(channel_id=space.channel.id, member_type="human", member_id=space.other.id))
        session.commit()
    assert space.api.get(path, headers=space.other_auth).status_code == 404
    publish(space, doc)
    assert space.api.get(path, headers=space.other_auth).status_code == 200
    assert space.api.delete(path, headers=space.other_auth).status_code == 403
    with db() as session:
        session.add(DirectMessage(channel_id=space.channel.id, user_id=space.user.id, agent_id=space.agent.id))
        session.commit()
    assert space.api.get(path, headers=space.other_auth).status_code == 404
    assert space.api.get(path + "/download", headers=space.auth).content.startswith(b"Project Amber")


def test_forged_foreign_attachment_cannot_create_a_task(space):
    doc = upload(space)
    with db() as session:
        item = session.get(Document, doc["id"])
        item.channel_id += 100000
        session.add(item)
        session.commit()
    response = space.api.post(f"/api/channels/{space.channel.id}/messages", headers=space.auth,
        json={"content": "@" + space.agent.name + " read", "attachment_ids": [doc["id"]]})
    assert response.status_code == 404
    with db() as session:
        assert not session.exec(select(Message).where(Message.channel_id == space.channel.id)).first()


@pytest.mark.parametrize("name,data", [("../bad.pdf", b"%PDF-"), ("broken.pdf", b"not pdf"),
    ("macro.xlsm", b"macro"), ("empty.txt", b""), ("image.png", b"not png")])
def test_rejects_invalid_uploads(space, name, data):
    response = space.api.post(f"/api/channels/{space.channel.id}/documents", params={"name": name},
                              content=data, headers=space.auth)
    assert response.status_code in {413, 422}


def test_word_and_spreadsheet_extraction_has_source_locations(space):
    from docx import Document as Word
    from openpyxl import Workbook
    word = Word()
    word.add_paragraph("Contract total 42")
    word.add_table(1, 2).rows[0].cells[0].text = "Amount"
    word.add_paragraph("Terms after the table")
    data = io.BytesIO()
    word.save(data)
    word_doc = upload(space, "contract.docx", data.getvalue())
    with db() as session:
        units = documents.extract(session.get(Document, word_doc["id"]))
    assert any(u["location"] == "Table 1" and "Amount" in u["text"] for u in units)
    assert [u["location"] for u in units] == ["Paragraph 1", "Table 1", "Paragraph 2"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Amount"])
    for i in range(1, 1501):
        sheet.append([i])
    data = io.BytesIO()
    workbook.save(data)
    doc = upload(space, "sales.xlsx", data.getvalue())
    with db() as session:
        units = documents.extract(session.get(Document, doc["id"]))
    assert "1500" in units[-1]["text"] and units[-1]["location"].startswith("Sales, rows")


def test_scan_ingestion_retries_reuse_file_ids_and_count_usage(space, monkeypatch):
    image = Image.new("RGB", (120, 100), "white")
    data = io.BytesIO()
    image.save(data, format="PDF")
    doc = upload(space, "scan.pdf", data.getvalue())
    publish(space, doc)
    fake = sdk_mock()
    monkeypatch.setattr(documents, "client", fake.client)
    asyncio.run(documents.ingest(doc["id"]))
    with db() as session:
        saved = session.get(Document, doc["id"])
        assert saved.status == "ready" and saved.units[0]["location"] == "Page 1"
        assert "1234" in saved.units[0]["text"]
        assert session.exec(select(DocumentUsage).where(DocumentUsage.channel_id == space.channel.id)).one().input_tokens == 30
    asyncio.run(documents.ingest(doc["id"]))
    assert fake.sdk.files.create.await_count == 2 and fake.sdk.responses.create.await_count == 1
    assert "file_id" not in space.api.get(f"/api/documents/{doc['id']}", headers=space.auth).json()


def test_search_suppresses_eventually_deleted_files_and_agent_revocation(space, monkeypatch):
    import app.runtime.tools.documents as tools
    doc = upload(space)
    publish(space, doc)
    documents.patch(doc["id"], status="ready", index_file_id="index")
    with db() as session:
        session.merge(DocumentCollection(channel_id=space.channel.id, vector_store_id="vector"))
        session.commit()
    fake = sdk_mock()
    fake.sdk.vector_stores.search.return_value = NS(data=[NS(attributes={"document_id": doc["id"]},
                                                           file_id="index", content=[NS(text="private result")])])
    monkeypatch.setattr(tools, "client", fake.client)
    search = {t.name: t for t in document_tools(space.task)}["search_documents"]
    assert asyncio.run(search.ainvoke({"query": "budget"}))["results"][0]["name"] == doc["name"]
    assert space.api.delete(f"/api/documents/{doc['id']}", headers=space.auth).status_code == 200
    assert asyncio.run(search.ainvoke({"query": "budget"}))["results"] == []
    with db() as session:
        session.delete(session.get(ChannelMember, (space.channel.id, "agent", space.agent.id)))
        session.commit()
    with pytest.raises(ValueError, match="access"):
        asyncio.run(search.ainvoke({"query": "budget"}))


@pytest.mark.parametrize("provider", ["openai", "codex", "claude"])
def test_calculation_receives_full_original_in_task_scoped_python_container(space, monkeypatch, provider):
    import app.runtime.tools.documents as tools
    doc = upload(space, "sales.csv", b"amount\n" + b"2\n" * 1500)
    publish(space, doc)
    documents.patch(doc["id"], status="ready", file_id="full-original")
    fake = sdk_mock()
    monkeypatch.setattr(tools, "client", fake.client)
    if provider != "openai":
        with db() as session:
            pinned = session.get(TaskModel, space.task.id)
            pinned.provider, pinned.model = provider, "claude-test" if provider == "claude" else "gpt-connected-account"
            session.add(pinned)
            session.commit()
    analyze = {t.name: t for t in document_tools(space.task)}["analyze_documents"]
    result = asyncio.run(analyze.ainvoke({"document_ids": [doc["id"]], "question": "Sum every row", "calculate": True}))
    call = fake.sdk.responses.create.call_args.kwargs
    if provider != "openai":
        assert call["model"] == configuration()[2]
    assert call["tools"] == [{"type": "code_interpreter", "container": {"type": "auto", "file_ids": ["full-original"]}}]
    assert result["complete"] and "full-original" not in str(result)


def test_ingestion_usage_blocks_agent_at_channel_budget(space):
    with db() as session:
        session.add(Budget(scope="channel", scope_id=space.channel.id, daily_tokens=40))
        session.add(DocumentUsage(id=uuid4().hex, channel_id=space.channel.id, model="gpt-6-astra", input_tokens=40))
        session.commit()
        assert over_budget(session, space.agent.id, space.channel.id)
    with pytest.raises(ValueError, match="budget"):
        documents.check_budget(space.channel.id)


@pytest.mark.parametrize("padding", ["", " ", "\t"])
@pytest.mark.parametrize("origin", ["", "https://chatgpt.com"])
def test_citations_require_a_document_read_and_are_revoked_on_delete(space, padding, origin):
    doc = upload(space)
    publish(space, doc)
    documents.patch(doc["id"], status="ready", units=[{"location": "Paragraph 1", "text": "Verified source"}])
    link = f"[source]({padding}{origin}/api/documents/{doc['id']}/download{padding})"
    assert "source unavailable" in documents.validate_citations(space.task, link)
    read = {t.name: t for t in document_tools(space.task)}["read_document"]
    assert read.invoke({"document_id": doc["id"]})["text"].endswith("Verified source")
    assert documents.validate_citations(space.task, link) == f"[source](/api/documents/{doc['id']}/download)"
    space.api.delete(f"/api/documents/{doc['id']}", headers=space.auth)
    assert "source unavailable" in documents.validate_citations(space.task, link)


def test_deleted_document_retains_in_flight_upload_id_for_cleanup(space):
    doc = upload(space)
    space.api.delete(f"/api/documents/{doc['id']}", headers=space.auth)
    with pytest.raises(ValueError, match="removed"):
        documents.patch(doc["id"], file_id="finished-in-flight")
    with db() as session:
        saved = session.get(Document, doc["id"])
        assert saved.status == "deleting" and saved.file_id == "finished-in-flight"


def test_completed_deletion_cannot_be_resurrected_by_late_worker_updates(space):
    doc = upload(space)
    space.api.delete(f"/api/documents/{doc['id']}", headers=space.auth)
    documents.patch(doc["id"], status="deleted", units=[], file_id="", index_file_id="")
    for values in ({"status": "ready"}, {"status": "error", "error": "late failure"}, {"file_id": "late-upload"}):
        with pytest.raises(ValueError, match="removed"):
            documents.patch(doc["id"], **values)
    with db() as session:
        saved = session.get(Document, doc["id"])
        assert saved.status == "deleted" and saved.file_id == "" and saved.units == []


def test_incomplete_scan_is_not_indexed_as_a_complete_document(space, monkeypatch):
    image = Image.new("RGB", (120, 100), "white")
    data = io.BytesIO()
    image.save(data, format="PDF")
    doc = upload(space, "scan.pdf", data.getvalue())
    publish(space, doc)
    fake = sdk_mock()
    fake.sdk.responses.create.return_value.status = "incomplete"
    monkeypatch.setattr(documents, "client", fake.client)
    assert asyncio.run(documents.process_one())
    saved = space.api.get(f"/api/documents/{doc['id']}", headers=space.auth).json()
    assert saved["status"] == "error" and "incomplete" in saved["error"]
    fake.sdk.files.create.assert_not_awaited()
    with db() as session:
        assert session.get(DocumentUsage, "response").output_tokens == 10


def test_csv_extraction_stops_at_the_row_limit(space):
    doc = upload(space, "too-many.csv", b"value\n" * 50001)
    with db() as session:
        with pytest.raises(ValueError, match="processing limit"):
            documents.extract(session.get(Document, doc["id"]))


@pytest.mark.parametrize("content,expected", [(b"amount\n42\n", True), (b"x" * 21, False), (b"", False)])
def test_generated_downloads_are_bounded_validated_and_closed(monkeypatch, content, expected):
    from app.runtime.tools.documents import generated_file
    monkeypatch.setattr(documents, "LIMIT", 20)
    class Response:
        closed = False
        async def __aenter__(self): return self
        async def __aexit__(self, *args): self.closed = True
        async def iter_bytes(self, chunk_size):
            for start in range(0, len(content), 5):
                yield content[start:start + 5]
    response = Response()
    sdk = NS(containers=NS(files=NS(content=NS(with_streaming_response=NS(retrieve=lambda **kwargs: response)))))
    annotation = NS(container_id="task-container", file_id="generated", filename="/mnt/data/summary.csv")
    result = asyncio.run(generated_file(sdk, annotation))
    assert (result is not None) == expected
    assert response.closed
    if expected:
        assert result == ("summary.csv", "text/csv", content)


def test_generated_download_is_linked_to_its_task_and_can_be_downloaded(space, monkeypatch):
    import app.runtime.tools.documents as tools
    from app.persistence.models import TaskDocument, TaskDocumentSource
    doc = upload(space, "sales.csv", b"amount\n42\n")
    publish(space, doc)
    documents.patch(doc["id"], status="ready", file_id="full-original")
    fake = sdk_mock()
    fake.sdk.responses.create.return_value.output = [NS(content=[NS(annotations=[NS(type="container_file_citation",
        container_id="task-container", file_id="summary", filename="summary.csv")])])]
    monkeypatch.setattr(tools, "client", fake.client)
    monkeypatch.setattr(tools, "generated_file", AsyncMock(return_value=("summary.csv", "text/csv", b"total\n42\n")))
    analyze = {t.name: t for t in document_tools(space.task)}["analyze_documents"]
    result = asyncio.run(analyze.ainvoke({"document_ids": [doc["id"]], "question": "Sum", "calculate": True}))
    download = result["downloads"][0]
    document_id = download["url"].split("/")[3]
    with db() as session:
        assert session.get(TaskDocument, (space.task.id, document_id))
        assert session.get(TaskDocumentSource, (space.task.id, document_id))
    assert space.api.get(download["url"], headers=space.auth).content == b"total\n42\n"


def test_openai_settings_are_admin_only_encrypted_and_masked(space, monkeypatch):
    import app.api.openai_settings as settings_api
    monkeypatch.setattr(settings_api, "validate_models", AsyncMock())
    url = "/api/settings/openai"
    assert space.api.get(url, headers=space.other_auth).status_code == 403
    body = {"api_key": "offline-private-key", "models": ["gpt-6-astra"], "default_model": "gpt-6-astra"}
    response = space.api.put(url, headers=space.auth, json=body)
    assert response.status_code == 200 and "offline-private-key" not in response.text
    with db() as session:
        assert session.get(OpenAISettings, 1).api_key_enc != "offline-private-key"
    assert configuration()[0] == "offline-private-key"
    assert space.api.put(url, headers=space.auth, json={**body, "default_model": "missing"}).status_code == 422


def test_model_snapshot_survives_profile_changes_and_legacy_restart_is_fresh(space):
    first = pin_task(space.task, space.agent)
    space.agent.model = "gpt-6-luna"
    assert pin_task(space.task, space.agent) == first
    with db() as session:
        task = session.get(AgentTask, space.task.id)
        session.delete(session.get(TaskModel, task.id))
        task.attempts, task.status = 2, "failed"
        agent = session.get(AgentProfile, space.agent.id)
        agent.model = "claude-sonnet-5-5"
        session.add(task)
        session.add(agent)
        session.commit()
    migrate_profiles()
    with pytest.raises(LegacyCheckpoint):
        pin_task(space.task, space.agent)
    response = space.api.post(f"/api/tasks/{space.task.id}/retry", headers=space.auth)
    assert response.status_code == 409
    response = space.api.post(f"/api/tasks/{space.task.id}/restart-with-providers", headers=space.auth)
    assert response.status_code == 200 and response.json()["id"] != space.task.id
